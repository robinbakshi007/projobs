from docx import Document

from document_engine.tailor_cv import _apply_rewrites


def test_apply_rewrites_replaces_summary_placeholder(tmp_path):
    p = tmp_path / "cv.docx"
    doc = Document()
    doc.add_paragraph("[SUMMARY] old summary")
    doc.save(p)

    loaded = Document(p)
    _apply_rewrites(loaded, {"[SUMMARY]": "new summary"})
    loaded.save(p)

    out = Document(p)
    assert out.paragraphs[0].text.startswith("[SUMMARY] new summary")
